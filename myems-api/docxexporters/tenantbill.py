"""
Tenant Bill DOCX Exporter

This module provides functionality to export the tenant bill (payment notice)
to DOCX format.

The exported DOCX file includes:
- Cover page with report metadata
- Payment notice section (logo, lease number, tenant contact block, bill header,
  billing detail table and the subtotal / VAT / total amount payable footer)

Note: the content mirrors the PDF exporter (pdfexporters/tenantbill.py)
exactly - the same columns, the same rows and the same 2-decimal precision
(3 decimals for the quantity column).
"""

import base64
import datetime
import os
import time
import uuid
import io

from decimal import Decimal
from typing import Optional, Dict, Any, BinaryIO
import logging

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from core.utilities import get_translation, round2

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

_font_setup_done = False


def setup_chinese_fonts():
    global _font_setup_done
    if _font_setup_done:
        return True

    font_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             '..', 'pdfexporters', 'fonts', 'NotoSansCJK-Regular.ttc')
    try:
        import matplotlib.font_manager as fm
        fm.fontManager.addfont(font_path)
        prop = fm.FontProperties(fname=font_path)
        font_name = prop.get_name()
        plt.rcParams['font.sans-serif'] = [font_name]
        plt.rcParams['axes.unicode_minus'] = False
        logger.info("Successfully loaded bundled font: %s from %s", font_name, font_path)
        _font_setup_done = True
        return True
    except Exception as e:
        logger.warning("Failed to load bundled font from %s: %s", font_path, e)

    plt.rcParams['font.sans-serif'] = ['DejaVu Sans']
    plt.rcParams['axes.unicode_minus'] = False
    logger.warning("Failed to load bundled NotoSansCJK font, using DejaVu Sans")
    return False


def _convert_decimals(obj):
    if isinstance(obj, Decimal):
        return float(obj)
    elif isinstance(obj, dict):
        return {k: _convert_decimals(v) for k, v in obj.items()}
    elif isinstance(obj, list):
        return [_convert_decimals(item) for item in obj]
    elif isinstance(obj, tuple):
        return tuple(_convert_decimals(item) for item in obj)
    return obj


def _to_date_str(datetime_str):
    if not datetime_str:
        return ''
    try:
        return datetime.datetime.strptime(datetime_str, '%Y-%m-%dT%H:%M:%S').isoformat()[0:10]
    except Exception:
        return str(datetime_str)[0:10]


def _set_min_row_height(cell, height_twips=150, exact=False):
    tr = cell._tc.getparent()
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(height_twips))
    trHeight.set(qn('w:hRule'), 'exact' if exact else 'atLeast')
    existing = trPr.find(qn('w:trHeight'))
    if existing is not None:
        trPr.remove(existing)
    trPr.append(trHeight)


def _reset_cell_paragraph_spacing(cell, font_size=9):
    line_twips = max(int(font_size * 20 * 1.15), 120)
    for p in cell.paragraphs:
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pPr = p._p.get_or_add_pPr()
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:before'), '0')
        spacing.set(qn('w:after'), '0')
        spacing.set(qn('w:line'), str(line_twips))
        spacing.set(qn('w:lineRule'), 'exact')
        existing = pPr.find(qn('w:spacing'))
        if existing is not None:
            pPr.remove(existing)
        pPr.append(spacing)
        ind = pPr.find(qn('w:ind'))
        if ind is not None:
            pPr.remove(ind)


def _set_cell_margins_zero(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side in ['top', 'start', 'bottom', 'end']:
        m = OxmlElement('w:' + side)
        m.set(qn('w:w'), '0')
        m.set(qn('w:type'), 'dxa')
        tcMar.append(m)
    existing = tcPr.find(qn('w:tcMar'))
    if existing is not None:
        tcPr.remove(existing)
    tcPr.append(tcMar)


def _remove_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement('w:' + border_name)
        border.set(qn('w:val'), 'none')
        border.set(qn('w:sz'), '0')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'auto')
        tblBorders.append(border)
    existing = tblPr.find(qn('w:tblBorders'))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(tblBorders)
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)


def _style_table_cell(cell, is_header=False, is_green=False, bold=False, font_size=9):
    if not cell.paragraphs[0].runs:
        cell.paragraphs[0].add_run('')
    cell.paragraphs[0].runs[0].font.bold = bold
    cell.paragraphs[0].runs[0].font.size = Pt(font_size)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _reset_cell_paragraph_spacing(cell, font_size=font_size)
    _set_cell_margins_zero(cell)
    row_h = max(int(font_size * 20 * 1.3), 160)
    _set_min_row_height(cell, height_twips=row_h, exact=False)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement('w:' + border_name)
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '666666')
        tcBorders.append(border)
    tcPr.append(tcBorders)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    if is_header:
        shd.set(qn('w:fill'), 'D9E2F3')
    elif is_green:
        shd.set(qn('w:fill'), '90EE90')
    else:
        return
    tcPr.append(shd)


class TenantBillDOCXExporter:
    """
    Export tenant bill (payment notice) data to DOCX format.
    """

    def __init__(self, language: str = 'zh_CN'):
        font_setup_success = setup_chinese_fonts()
        if not font_setup_success:
            logger.warning("Chinese font setup failed, some text may not display correctly")

        self.language = language
        self.trans = get_translation(language)
        self._ = self.trans.gettext

        self.dpi = 120

    def export(self,
               report: Dict[str, Any],
               name: str,
               reporting_start_datetime_local: str,
               reporting_end_datetime_local: str,
               period_type: str,
               language: str) -> Optional[str]:
        if report is None:
            return None
        start_time = time.time()
        logger.info("Starting DOCX generation for tenant bill: %s", name)

        docx_filename = self.generate_docx(
            report, name,
            reporting_start_datetime_local,
            reporting_end_datetime_local,
            period_type,
            language
        )

        result = ''
        if docx_filename and os.path.exists(docx_filename):
            try:
                with open(docx_filename, 'rb') as binary_file:
                    binary_data = binary_file.read()
                result = base64.b64encode(binary_data).decode('utf-8')
            except Exception as e:
                logger.error("Failed to encode DOCX: %s", str(e))
            finally:
                try:
                    os.remove(docx_filename)
                except Exception:
                    pass
        elapsed = time.time() - start_time
        logger.info("DOCX generation completed in %.2fs for tenant bill: %s", elapsed, name)
        return result

    @staticmethod
    def _fig_to_bytesio(fig, dpi: int) -> BinaryIO:
        buf = io.BytesIO()
        fig.tight_layout()
        fig.savefig(buf, format='png', dpi=dpi, bbox_inches='tight')
        plt.close(fig)
        buf.seek(0)
        return buf

    def generate_docx(self,
                      report: Dict[str, Any],
                      name: str,
                      reporting_start_datetime_local: str,
                      reporting_end_datetime_local: str,
                      period_type: str,
                      language: str) -> Optional[str]:
        _ = self._

        filename = str(uuid.uuid4()) + '.docx'

        self.report = _convert_decimals(report)
        self.name = name
        self.reporting_start = reporting_start_datetime_local
        self.reporting_end = reporting_end_datetime_local
        self.period_type = 'daily'

        doc = Document()
        section = doc.sections[0]
        section.orientation = 1
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)

        self._add_cover_page(doc, name,
                             reporting_start_datetime_local,
                             reporting_end_datetime_local)

        self._add_payment_notice_section(doc)

        doc.save(filename)
        logger.info("DOCX generated: %s", filename)
        return filename

    def _add_heading_styled(self, doc, text, level=1):
        heading = doc.add_heading(level=level)
        run = heading.add_run(text)
        run.font.bold = True
        run.font.name = 'Arial'
        r = run._element
        r.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
        return heading

    def _add_cover_page(self, doc, name, reporting_start, reporting_end):
        _ = self._
        img_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                '..', 'excelexporters', 'myems.png')
        if os.path.exists(img_path):
            try:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(img_path, width=Inches(7.0))
            except Exception as e:
                logger.warning("Failed to load logo image: %s", e)

        for _unused in range(3):
            doc.add_paragraph('')

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(_('Tenant Data') + ' - ' + _('Tenant Bill'))
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.name = 'Arial'
        r = run._element
        r.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

        for _unused in range(3):
            doc.add_paragraph('')

        info_data = [
            [_('Tenant') + ':', name],
            [_('Reporting Start Datetime') + ':', reporting_start],
            [_('Reporting End Datetime') + ':', reporting_end],
        ]

        info_table = doc.add_table(rows=len(info_data), cols=2)
        info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _remove_table_borders(info_table)

        half_w = Inches(3.2)
        for i, (label, value) in enumerate(info_data):
            c_label = info_table.cell(i, 0)
            c_label.width = half_w
            c_label.text = label
            c_label.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_lbl = c_label.paragraphs[0].runs[0]
            r_lbl.bold = True
            r_lbl.font.size = Pt(13)
            r_lbl.font.name = 'Arial'
            r_lbl._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

            c_value = info_table.cell(i, 1)
            c_value.width = half_w
            c_value.text = str(value) if value is not None else ''
            c_value.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_val = c_value.paragraphs[0].runs[0]
            r_val.font.size = Pt(13)
            r_val.font.name = 'Arial'
            r_val._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

        doc.add_page_break()

    def _build_detail_rows(self):
        _ = self._
        reporting_period = self.report.get('reporting_period') or {}
        names = reporting_period.get('names') or []
        units = reporting_period.get('units') or []
        subtotals_input = reporting_period.get('subtotals_input') or []
        subtotals_cost = reporting_period.get('subtotals_cost') or []

        start_date = _to_date_str(self.reporting_start)
        end_date = _to_date_str(self.reporting_end)

        col_headers = [_('Energy Category'), _('Billing Period Start'), _('Billing Period End'),
                       _('Quantity'), _('Unit'), _('Amount'), _('Tax Rate'), _('VAT Output Tax')]

        rows = []
        for i in range(len(names)):
            quantity = subtotals_input[i] if i < len(subtotals_input) else None
            amount = subtotals_cost[i] if i < len(subtotals_cost) else None
            rows.append([
                str(names[i]),
                start_date,
                end_date,
                str(round2(quantity, 3)) if quantity is not None else '',
                str(units[i]) if i < len(units) else '',
                str(round2(amount, 2)) if amount is not None else '',
                '0',
                '0',
            ])
        return col_headers, rows

    def _add_payment_notice_section(self, doc):
        _ = self._

        tenant_data = self.report.get('tenant') or {}
        reporting_period = self.report.get('reporting_period') or {}
        currency_unit = reporting_period.get('currency_unit') or ''
        total_cost = reporting_period.get('total_cost')
        total_cost = round2(total_cost, 2) if total_cost is not None else round2(0, 2)
        taxes = round2(0, 2)

        col_headers, detail_rows = self._build_detail_rows()
        total_cols = len(col_headers)

        first_page_rows = 50
        other_page_rows = 60

        pages = []
        if not detail_rows:
            pages.append([])
        else:
            pages.append(detail_rows[0:first_page_rows])
            offset = first_page_rows
            while offset < len(detail_rows):
                pages.append(detail_rows[offset:offset + other_page_rows])
                offset += other_page_rows

        page_total = len(pages)
        for page_index, page_rows in enumerate(pages):
            notice_container = doc.add_table(rows=3, cols=1)
            notice_container.alignment = WD_TABLE_ALIGNMENT.CENTER
            _remove_table_borders(notice_container)

            header_cell = notice_container.cell(0, 0)
            if header_cell.paragraphs:
                p = header_cell.paragraphs[0]._element
                p.getparent().remove(p)
            if page_index == 0:
                self._draw_notice_header(header_cell, tenant_data, currency_unit, total_cost, taxes)
            else:
                self._draw_notice_simple_title(header_cell)

            table_cell = notice_container.cell(1, 0)
            if table_cell.paragraphs:
                p = table_cell.paragraphs[0]._element
                p.getparent().remove(p)
            for _unused in range(2):
                p = table_cell.add_paragraph('')
                pf = p.paragraph_format
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            table_data = [col_headers] + page_rows
            num_data_rows = len(table_data)

            data_table = table_cell.add_table(rows=num_data_rows, cols=total_cols)
            data_table.alignment = WD_TABLE_ALIGNMENT.CENTER

            table_font_size = 7
            for j_col in range(total_cols):
                cell = data_table.cell(0, j_col)
                cell.text = col_headers[j_col]
                _style_table_cell(cell, is_green=True, bold=True, font_size=table_font_size)

            for i_row in range(1, num_data_rows):
                for j_col in range(total_cols):
                    cell = data_table.cell(i_row, j_col)
                    cell.text = table_data[i_row][j_col]
                    _style_table_cell(cell, font_size=table_font_size)

            for _unused in range(2):
                p = table_cell.add_paragraph('')
                pf = p.paragraph_format
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

            footer_cell = notice_container.cell(2, 0)
            if footer_cell.paragraphs:
                p = footer_cell.paragraphs[0]._element
                p.getparent().remove(p)
            is_last = page_index == page_total - 1
            self._draw_notice_footer(footer_cell, currency_unit, total_cost, taxes, is_last)

            if page_index < page_total - 1:
                doc.add_page_break()

    @staticmethod
    def _add_floating_logo_top_right(paragraph, img_path, logo_height_inches=1.0):
        """Insert logo as a floating shape anchored to the given paragraph.

        The image is positioned with:
          - Text wrapping: Behind text (behindDoc="1")
          - Locked anchor so it stays fixed on page
          - Horizontally aligned to the right of the page margin
          - Vertically aligned to the top of the page margin
        Returns True on success, False otherwise.
        """
        try:
            from docx.oxml.ns import nsmap
            run = paragraph.add_run()
            run.add_picture(img_path, height=Inches(logo_height_inches))

            drawing = run._element.find(qn('w:drawing'))
            if drawing is None:
                return False
            inline = drawing.find(qn('wp:inline'))
            if inline is None:
                return False

            cx = inline.get('distT') or '0'
            cy = inline.get('distB') or '0'
            cx_l = inline.get('distL') or '0'
            cx_r = inline.get('distR') or '0'
            extent = inline.find(qn('wp:extent'))
            effect_extent = inline.find(qn('wp:effectExtent'))
            doc_pr = inline.find(qn('wp:docPr'))
            c_nv_gp_frame_pr = inline.find(qn('wp:cNvGraphicFramePr'))
            graphic = inline.find(qn('a:graphic'))

            drawing.remove(inline)

            anchor = OxmlElement('wp:anchor')
            anchor.set('distT', cx)
            anchor.set('distB', cy)
            anchor.set('distL', cx_l)
            anchor.set('distR', cx_r)
            anchor.set('simplePos', '0')
            anchor.set('relativeHeight', '251658240')
            anchor.set('behindDoc', '1')
            anchor.set('locked', '1')
            anchor.set('layoutInCell', '1')
            anchor.set('allowOverlap', '1')

            simple_pos = OxmlElement('wp:simplePos')
            simple_pos.set('x', '0')
            simple_pos.set('y', '0')
            anchor.append(simple_pos)

            h_pos = OxmlElement('wp:positionH')
            h_pos.set('relativeFrom', 'margin')
            h_align = OxmlElement('wp:align')
            h_align.text = 'right'
            h_pos.append(h_align)
            anchor.append(h_pos)

            v_pos = OxmlElement('wp:positionV')
            v_pos.set('relativeFrom', 'margin')
            v_align = OxmlElement('wp:align')
            v_align.text = 'top'
            v_pos.append(v_align)
            anchor.append(v_pos)

            wrap_none = OxmlElement('wp:wrapNone')
            anchor.append(wrap_none)

            if extent is not None:
                anchor.append(extent)
            if effect_extent is not None:
                anchor.append(effect_extent)
            if doc_pr is not None:
                anchor.append(doc_pr)
            if c_nv_gp_frame_pr is not None:
                anchor.append(c_nv_gp_frame_pr)
            if graphic is not None:
                anchor.append(graphic)

            drawing.append(anchor)
            return True
        except Exception as e:
            logger.warning("Failed to insert floating logo: %s", e)
            return False

    def _draw_notice_header(self, container_cell, tenant_data, currency_unit, total_cost, taxes):
        _ = self._

        lease_number = tenant_data.get('lease_number')

        header_info_container = container_cell.add_table(rows=2, cols=1)
        _remove_table_borders(header_info_container)

        title_row = header_info_container.cell(0, 0)
        if title_row.paragraphs:
            p = title_row.paragraphs[0]._element
            p.getparent().remove(p)
        p_logo_anchor = title_row.add_paragraph()

        img_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                '..', 'excelexporters', 'myemslogo.png')
        if os.path.exists(img_path):
            self._add_floating_logo_top_right(p_logo_anchor, img_path, logo_height_inches=1.0)

        p_title = title_row.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run(_('Payment Notice'))
        run_title.font.size = Pt(20)
        run_title.font.bold = True
        run_title.font.name = 'Arial'
        r_title = run_title._element
        r_title.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

        if lease_number:
            p_lease = title_row.add_paragraph()
            p_lease.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run_lease = p_lease.add_run(_('Lease Contract Number') + ': ' + str(lease_number))
            run_lease.font.size = Pt(11)
            run_lease.font.bold = True
            run_lease.font.name = 'Arial'
            r_lease = run_lease._element
            r_lease.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

        contact_row = header_info_container.cell(1, 0)
        if contact_row.paragraphs:
            p = contact_row.paragraphs[0]._element
            p.getparent().remove(p)
        contact_container = contact_row.add_table(rows=1, cols=2)
        contact_container.alignment = WD_TABLE_ALIGNMENT.CENTER
        _remove_table_borders(contact_container)

        left_cell = contact_container.cell(0, 0)
        if left_cell.paragraphs:
            p = left_cell.paragraphs[0]._element
            p.getparent().remove(p)
        right_cell = contact_container.cell(0, 1)
        if right_cell.paragraphs:
            p = right_cell.paragraphs[0]._element
            p.getparent().remove(p)

        left_rows = [
            str(tenant_data.get('name', '') or ''),
            str(tenant_data.get('rooms', '') or ''),
            str(tenant_data.get('floors', '') or ''),
            str(tenant_data.get('buildings', '') or ''),
            str(tenant_data.get('email', '') or ''),
            str(tenant_data.get('phone', '') or ''),
        ]
        for item in left_rows:
            p = left_cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(item)
            run.font.size = Pt(9)
            run.font.name = 'Arial'
            r = run._element
            r.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

        right_rows = [
            (_('Bill Number'), ''),
            (_('Lease Contract Number'), str(lease_number or '')),
            (_('Bill Date'), _to_date_str(self.reporting_start)),
            (_('Payment Due Date'), _to_date_str(self.reporting_end)),
            (_('Amount Payable'), currency_unit + str(total_cost)),
        ]
        right_table = right_cell.add_table(rows=len(right_rows), cols=2)
        _remove_table_borders(right_table)
        for i, (label, value) in enumerate(right_rows):
            c_lbl = right_table.cell(i, 0)
            c_lbl.text = label
            c_lbl.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_lbl = c_lbl.paragraphs[0].runs[0]
            r_lbl.bold = True
            r_lbl.font.size = Pt(9)
            r_lbl.font.name = 'Arial'
            r_lbl._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

            c_val = right_table.cell(i, 1)
            c_val.text = str(value) if value is not None else ''
            c_val.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_val = c_val.paragraphs[0].runs[0]
            r_val.font.size = Pt(9)
            r_val.font.name = 'Arial'
            r_val._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

    def _draw_notice_simple_title(self, container_cell):
        _ = self._
        p_title = container_cell.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run(_('Payment Notice'))
        run_title.font.size = Pt(16)
        run_title.font.bold = True
        run_title.font.name = 'Arial'
        r_title = run_title._element
        r_title.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

        p_sub = container_cell.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sub = p_sub.add_run(self.name + ' (' + _('Tenant Bill') + ')')
        run_sub.font.size = Pt(9)
        run_sub.font.name = 'Arial'
        r_sub = run_sub._element
        r_sub.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

    def _draw_notice_footer(self, container_cell, currency_unit, total_cost, taxes, is_last=True):
        _ = self._

        if not is_last:
            return

        footer_rows = [
            [_('Subtotal') + ':', currency_unit + str(total_cost)],
            [_('VAT Output Tax') + ':', currency_unit + str(taxes)],
            [_('Total Amount Payable') + ':', currency_unit + str(round2(total_cost + taxes, 2))],
        ]

        footer_container = container_cell.add_table(rows=1, cols=2)
        footer_container.alignment = WD_TABLE_ALIGNMENT.CENTER
        _remove_table_borders(footer_container)
        footer_container.cell(0, 0).merge(footer_container.cell(0, 1))
        merged_cell = footer_container.cell(0, 0)
        if merged_cell.paragraphs:
            p = merged_cell.paragraphs[0]._element
            p.getparent().remove(p)

        right_align_container = merged_cell.add_table(rows=1, cols=3)
        _remove_table_borders(right_align_container)
        right_cell = right_align_container.cell(0, 2)
        if right_cell.paragraphs:
            p = right_cell.paragraphs[0]._element
            p.getparent().remove(p)

        inner_footer_table = right_cell.add_table(rows=len(footer_rows), cols=2)
        _remove_table_borders(inner_footer_table)

        for i, (label, value) in enumerate(footer_rows):
            c_lbl = inner_footer_table.cell(i, 0)
            c_lbl.text = label
            c_lbl.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_lbl = c_lbl.paragraphs[0].runs[0]
            r_lbl.font.size = Pt(10)
            r_lbl.font.name = 'Arial'
            r_lbl._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

            c_val = inner_footer_table.cell(i, 1)
            c_val.text = value
            c_val.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_val = c_val.paragraphs[0].runs[0]
            r_val.font.bold = True
            r_val.font.size = Pt(10)
            r_val.font.name = 'Arial'
            r_val._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')


def export(report, name, reporting_start_datetime_local,
           reporting_end_datetime_local, period_type, language):
    """
    Export report data to DOCX and return base64 encoded string.
    This function maintains the same interface as the PDF exporter.
    """
    exporter = TenantBillDOCXExporter(language)
    return exporter.export(report, name,
                           reporting_start_datetime_local,
                           reporting_end_datetime_local,
                           period_type,
                           language)
