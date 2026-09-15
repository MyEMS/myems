"""
Meter Comparison DOCX Exporter

This module provides functionality to export meter comparison data to DOCX format.
It generates comprehensive reports comparing energy consumption between two meters
with detailed breakdown by time periods.

Key Features:
- Two-meter energy consumption comparison
- Consumption summary table (each meter keeps its own energy category and unit)
- Detailed data table with timestamps, values, and difference
- Comparison line chart (both meters overlaid)
- Multi-language support
- Base64 encoding for file transmission

The exported DOCX file includes:
- Cover page with logo and report metadata
- Consumption summary table (meter1, meter2)
- Line chart of reporting period consumption (meter1, meter2)
- Detailed data table (paginated, matching Excel content)
"""

import base64
import os
import time
import uuid
import io

from decimal import Decimal
from typing import Optional, Dict, List, Any, BinaryIO
import logging

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from core.utilities import get_translation, round2

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

_font_setup_done = False


def setup_chinese_fonts():
    global _font_setup_done
    if _font_setup_done:
        return True

    font_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             '..', 'pdfexporters', 'fonts', 'NotoSansCJK-Regular.ttc')
    try:
        import matplotlib.font_manager as fm
        fm.fontManager.addfont(font_path)
        prop = fm.FontProperties(fname=font_path)
        font_name = prop.get_name()
        plt.rcParams['font.sans-serif'] = [font_name]
        plt.rcParams['axes.unicode_minus'] = False
        logger.info(f"Successfully loaded bundled font: {font_name} from {font_path}")
        _font_setup_done = True
        return True
    except Exception as e:
        logger.warning(f"Failed to load bundled font from {font_path}: {e}")

    plt.rcParams['font.sans-serif'] = ['DejaVu Sans']
    plt.rcParams['axes.unicode_minus'] = False
    logger.warning("Failed to load bundled NotoSansCJK font, using DejaVu Sans")
    return False


def _convert_decimals(obj):
    if isinstance(obj, Decimal):
        return float(obj)
    elif isinstance(obj, dict):
        return {k: _convert_decimals(v) for k, v in obj.items()}
    elif isinstance(obj, list):
        return [_convert_decimals(item) for item in obj]
    elif isinstance(obj, tuple):
        return tuple(_convert_decimals(item) for item in obj)
    return obj


def _set_min_row_height(cell, height_twips=150, exact=False):
    tr = cell._tc.getparent()
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(height_twips))
    trHeight.set(qn('w:hRule'), 'exact' if exact else 'atLeast')
    existing = trPr.find(qn('w:trHeight'))
    if existing is not None:
        trPr.remove(existing)
    trPr.append(trHeight)


def _reset_cell_paragraph_spacing(cell, font_size=9):
    line_twips = max(int(font_size * 20 * 1.15), 120)
    for p in cell.paragraphs:
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pPr = p._p.get_or_add_pPr()
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:before'), '0')
        spacing.set(qn('w:after'), '0')
        spacing.set(qn('w:line'), str(line_twips))
        spacing.set(qn('w:lineRule'), 'exact')
        existing = pPr.find(qn('w:spacing'))
        if existing is not None:
            pPr.remove(existing)
        pPr.append(spacing)
        ind = pPr.find(qn('w:ind'))
        if ind is not None:
            pPr.remove(ind)


def _set_cell_margins_zero(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side in ['top', 'start', 'bottom', 'end']:
        m = OxmlElement(f'w:{side}')
        m.set(qn('w:w'), '0')
        m.set(qn('w:type'), 'dxa')
        tcMar.append(m)
    existing = tcPr.find(qn('w:tcMar'))
    if existing is not None:
        tcPr.remove(existing)
    tcPr.append(tcMar)


def _remove_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        border.set(qn('w:sz'), '0')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'auto')
        tblBorders.append(border)
    existing = tblPr.find(qn('w:tblBorders'))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(tblBorders)
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)


def _style_table_cell(cell, is_header=False, is_green=False, bold=False, font_size=9):
    cell.paragraphs[0].runs[0].font.bold = bold
    cell.paragraphs[0].runs[0].font.size = Pt(font_size)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _reset_cell_paragraph_spacing(cell, font_size=font_size)
    _set_cell_margins_zero(cell)
    row_h = max(int(font_size * 20 * 1.3), 160)
    _set_min_row_height(cell, height_twips=row_h, exact=False)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '666666')
        tcBorders.append(border)
    tcPr.append(tcBorders)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    if is_header:
        shd.set(qn('w:fill'), 'D9E2F3')
    elif is_green:
        shd.set(qn('w:fill'), '90EE90')
    else:
        return
    tcPr.append(shd)


class MeterComparisonDOCXExporter:
    """
    Export meter comparison data to DOCX format.
    Generates comprehensive reports comparing two meters' energy consumption.
    Each meter keeps its own energy category name and unit of measure.
    """

    def __init__(self, language: str = 'zh_CN'):
        font_setup_success = setup_chinese_fonts()
        if not font_setup_success:
            logger.warning("Chinese font setup failed, some text may not display correctly")

        self.language = language
        self.trans = get_translation(language)
        self._ = self.trans.gettext

        self.dpi = 120

        self.chart_colors = ['#4472C4', '#ED7D31', '#70AD47', '#FFC000', '#5B9BD5',
                             '#FF6B6B', '#9B59B6', '#1ABC9C', '#E67E22', '#2ECC71',
                             '#3498DB', '#E74C3C', '#2ECC71', '#F39C12', '#9B59B6']

    def export(self,
               report: Dict[str, Any],
               name1: str,
               name2: str,
               reporting_start_datetime_local: str,
               reporting_end_datetime_local: str,
               period_type: str,
               language: str) -> Optional[str]:
        if report is None:
            return None
        start_time = time.time()
        logger.info(f"Starting DOCX generation for MeterComparison: {name1} vs {name2}")

        docx_filename = self.generate_docx(
            report, name1, name2,
            reporting_start_datetime_local, reporting_end_datetime_local,
            period_type, language
        )

        result = ''
        if docx_filename and os.path.exists(docx_filename):
            try:
                with open(docx_filename, 'rb') as binary_file:
                    binary_data = binary_file.read()
                result = base64.b64encode(binary_data).decode('utf-8')
            except Exception as e:
                logger.error(f"Failed to encode DOCX: {str(e)}")
            finally:
                try:
                    os.remove(docx_filename)
                except Exception:
                    pass
        elapsed = time.time() - start_time
        logger.info(f"DOCX generation completed in {elapsed:.2f}s for MeterComparison")
        return result

    @staticmethod
    def _fig_to_bytesio(fig, dpi: int) -> BinaryIO:
        buf = io.BytesIO()
        fig.tight_layout()
        fig.savefig(buf, format='png', dpi=dpi, bbox_inches='tight')
        plt.close(fig)
        buf.seek(0)
        return buf

    @staticmethod
    def _sanitize_values(values):
        result = []
        for v in values:
            if v is None:
                result.append(float('nan'))
            else:
                try:
                    result.append(float(v))
                except (TypeError, ValueError):
                    result.append(float('nan'))
        return result

    def generate_docx(self,
                      report: Dict[str, Any],
                      name1: str,
                      name2: str,
                      reporting_start_datetime_local: str,
                      reporting_end_datetime_local: str,
                      period_type: str,
                      language: str) -> Optional[str]:
        _ = self._

        if "reporting_period1" not in report.keys() or \
                "values" not in report['reporting_period1'].keys() or \
                len(report['reporting_period1']['values']) == 0:
            doc = Document()
            section = doc.sections[0]
            section.orientation = 1
            section.page_width = Inches(11.69)
            section.page_height = Inches(8.27)
            self._add_cover_page(doc, name1, name2,
                                 reporting_start_datetime_local,
                                 reporting_end_datetime_local,
                                 period_type,
                                 has_next=False)
            filename = str(uuid.uuid4()) + '.docx'
            doc.save(filename)
            return filename

        filename = str(uuid.uuid4()) + '.docx'

        self.report = _convert_decimals(report)
        self.name1 = name1
        self.name2 = name2
        self.reporting_start = reporting_start_datetime_local
        self.reporting_end = reporting_end_datetime_local
        self.period_type = period_type

        meter1 = self.report.get('meter1', {})
        meter2 = self.report.get('meter2', {})
        self.ec1 = meter1.get('energy_category_name', '')
        self.unit1 = meter1.get('unit_of_measure', '')
        self.ec2 = meter2.get('energy_category_name', '')
        self.unit2 = meter2.get('unit_of_measure', '')

        doc = Document()
        section = doc.sections[0]
        section.orientation = 1
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)

        rp1 = self.report.get('reporting_period1', {})
        rp1_timestamps = rp1.get('timestamps', [])
        has_summary = True
        has_chart = len(rp1_timestamps) > 0
        has_detailed = len(rp1_timestamps) > 0

        self._add_cover_page(doc, name1, name2,
                             reporting_start_datetime_local,
                             reporting_end_datetime_local,
                             period_type,
                             has_next=(has_summary or has_chart or has_detailed))

        self._add_consumption_summary(doc, has_next=(has_chart or has_detailed))
        self._add_line_chart_section(doc, has_next=has_detailed)
        self._add_detailed_data_section(doc)

        doc.save(filename)
        logger.info(f"DOCX generated: {filename}")
        return filename

    def _add_heading_styled(self, doc, text, level=1):
        heading = doc.add_heading(level=level)
        run = heading.add_run(text)
        run.font.bold = True
        run.font.name = 'Arial'
        r = run._element
        r.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
        return heading

    def _add_cover_page(self, doc, name1, name2,
                        reporting_start, reporting_end,
                        period_type,
                        has_next=True):
        _ = self._
        img_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                '..', 'excelexporters', 'myems.png')
        if os.path.exists(img_path):
            try:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(img_path, width=Inches(7.0))
            except Exception as e:
                logger.warning(f"Failed to load logo image: {e}")

        for _unused in range(2):
            doc.add_paragraph('')

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(_('Meter Data') + ' - ' + _('Meter Comparison'))
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.name = 'Arial'
        r = run._element
        r.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

        for _unused in range(2):
            doc.add_paragraph('')

        info_data = [
            [_('Name') + '1:', name1],
            [_('Name') + '2:', name2],
            [_('Period Type') + ':', period_type],
            [_('Reporting Start Datetime') + ':', reporting_start],
            [_('Reporting End Datetime') + ':', reporting_end],
        ]

        info_table = doc.add_table(rows=len(info_data), cols=2)
        info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _remove_table_borders(info_table)

        half_w = Inches(3.2)
        for i, (label, value) in enumerate(info_data):
            c_label = info_table.cell(i, 0)
            c_label.width = half_w
            c_label.text = label
            c_label.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_lbl = c_label.paragraphs[0].runs[0]
            r_lbl.bold = True
            r_lbl.font.size = Pt(13)
            r_lbl.font.name = 'Arial'
            r_lbl._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

            c_value = info_table.cell(i, 1)
            c_value.width = half_w
            c_value.text = str(value) if value is not None else ''
            c_value.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_val = c_value.paragraphs[0].runs[0]
            r_val.font.size = Pt(13)
            r_val.font.name = 'Arial'
            r_val._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

        if has_next:
            doc.add_page_break()

    def _add_consumption_summary(self, doc, has_next=True):
        """Add consumption summary table matching PDF layout:
        Two separate tables - Meter1 name/Consumption with EC1(unit1) and total1,
        Meter2 name/Consumption with EC2(unit2) and total2.
        Each meter keeps its own energy category name and unit.
        """
        _ = self._
        ec1_unit_label = self.ec1 + ' (' + self.unit1 + ')'
        ec2_unit_label = self.ec2 + ' (' + self.unit2 + ')'

        rp1 = self.report.get('reporting_period1', {})
        rp2 = self.report.get('reporting_period2', {})

        total1 = round2(rp1.get('total_in_category', 0), 2)
        total2 = round2(rp2.get('total_in_category', 0), 2)

        self._add_heading_styled(doc, self.name1 + ' & ' + self.name2 + ' - ' +
                                 _('Consumption'), level=1)

        container = doc.add_table(rows=2, cols=1)
        container.alignment = WD_TABLE_ALIGNMENT.CENTER
        _remove_table_borders(container)

        for meter_idx, (meter_name, ec_unit_label, total_val) in enumerate([
                (self.name1, ec1_unit_label, total1),
                (self.name2, ec2_unit_label, total2)]):
            row_cell = container.cell(meter_idx, 0)

            meter_table = row_cell.add_table(rows=2, cols=2)
            meter_table.alignment = WD_TABLE_ALIGNMENT.CENTER

            c00 = meter_table.cell(0, 0)
            c00.text = meter_name
            _style_table_cell(c00, is_green=True, bold=True)

            c01 = meter_table.cell(0, 1)
            c01.text = ec_unit_label
            _style_table_cell(c01, is_green=True, bold=True)

            c10 = meter_table.cell(1, 0)
            c10.text = _('Consumption')
            _style_table_cell(c10, is_green=True, bold=True)

            c11 = meter_table.cell(1, 1)
            c11.text = str(total_val)
            _style_table_cell(c11)

        if has_next:
            doc.add_paragraph('')
            doc.add_page_break()

    def _add_line_chart_section(self, doc, has_next=True):
        """Add line chart section matching PDF's 'Reporting Period Consumption' chart."""
        _ = self._
        rp1 = self.report.get('reporting_period1', {})
        rp2 = self.report.get('reporting_period2', {})

        timestamps = rp1.get('timestamps', [])
        values1 = rp1.get('values', [])
        values2 = rp2.get('values', [])

        if not timestamps or len(timestamps) == 0:
            return

        self._add_heading_styled(doc, self.name1 + ' & ' + self.name2 + ' - ' +
                                 _('Reporting Period Consumption'), level=1)

        xs = list(range(len(timestamps)))
        ys1 = self._sanitize_values(values1)
        ys2 = self._sanitize_values(values2)

        fig, ax = plt.subplots(figsize=(9.75, 5.25))
        marker_step = max(1, len(xs) // 30)

        m1_label = self.name1 + ' ' + self.ec1 + ' (' + self.unit1 + ')'
        m2_label = self.name2 + ' ' + self.ec2 + ' (' + self.unit2 + ')'

        valid_pairs1 = [(x, y) for x, y in zip(xs, ys1) if not (isinstance(y, float) and y != y)]
        if valid_pairs1:
            vx1, vy1 = zip(*valid_pairs1)
            ax.plot(vx1, vy1, marker='o', markersize=3,
                    linewidth=1.5, markevery=max(1, len(valid_pairs1) // 30),
                    label=m1_label, color='#4472C4')

        valid_pairs2 = [(x, y) for x, y in zip(xs, ys2) if not (isinstance(y, float) and y != y)]
        if valid_pairs2:
            vx2, vy2 = zip(*valid_pairs2)
            ax.plot(vx2, vy2, marker='s', markersize=3,
                    linewidth=1.5, markevery=max(1, len(valid_pairs2) // 30),
                    label=m2_label, color='#ED7D31')

        step = max(1, len(timestamps) // 10)
        ax.set_xticks(range(0, len(timestamps), step))
        ax.set_xticklabels(
            [timestamps[t][:10] for t in range(0, len(timestamps), step)],
            rotation=45, ha='right', fontsize=7)

        ax.set_ylabel(_('Consumption'), fontsize=9)
        ax.legend(fontsize=9)
        ax.grid(True, alpha=0.3)

        chart_buf = self._fig_to_bytesio(fig, self.dpi)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(chart_buf, width=Inches(9.75))

        if has_next:
            doc.add_page_break()

    def _add_detailed_data_section(self, doc):
        """Add detailed time-series data tables with timestamps, values, and difference."""
        _ = self._

        rp1 = self.report.get('reporting_period1', {})
        rp2 = self.report.get('reporting_period2', {})
        diff = self.report.get('diff', {})

        timestamps = rp1.get('timestamps', [])
        values1 = rp1.get('values', [])
        values2 = rp2.get('values', [])
        diff_values = diff.get('values', [])

        if not timestamps or len(timestamps) == 0:
            return

        rows_per_page = 30

        self._add_heading_styled(doc, self.name1 + ' and ' + self.name2 + ' ' +
                                 _('Detailed Data'), level=1)

        col_headers = [
            _('Datetime'),
            self.name1 + ' ' + self.ec1 + ' (' + self.unit1 + ')',
            self.name2 + ' ' + self.ec2 + ' (' + self.unit2 + ')',
            _('Difference')
        ]
        total_cols = len(col_headers)

        num_pages = (len(timestamps) + rows_per_page - 1) // rows_per_page

        for page in range(num_pages):
            start_row = page * rows_per_page
            end_row = min(start_row + rows_per_page, len(timestamps))
            page_rows = end_row - start_row

            table = doc.add_table(rows=page_rows + 1, cols=total_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for j, h in enumerate(col_headers):
                c = table.cell(0, j)
                c.text = h
                _style_table_cell(c, is_header=True, bold=True, font_size=8)

            for t_idx in range(page_rows):
                global_idx = start_row + t_idx
                r_idx = t_idx + 1
                c0 = table.cell(r_idx, 0)
                c0.text = str(timestamps[global_idx])
                _style_table_cell(c0, font_size=8)

                v1 = round2(values1[global_idx], 2) \
                    if global_idx < len(values1) and values1[global_idx] is not None else ''
                c1 = table.cell(r_idx, 1)
                c1.text = str(v1) if v1 != '' else ''
                _style_table_cell(c1, font_size=8)

                v2 = round2(values2[global_idx], 2) \
                    if global_idx < len(values2) and values2[global_idx] is not None else ''
                c2 = table.cell(r_idx, 2)
                c2.text = str(v2) if v2 != '' else ''
                _style_table_cell(c2, font_size=8)

                vd = round2(diff_values[global_idx], 2) \
                    if global_idx < len(diff_values) and diff_values[global_idx] is not None else ''
                c3 = table.cell(r_idx, 3)
                c3.text = str(vd) if vd != '' else ''
                _style_table_cell(c3, font_size=8)

            if page < num_pages - 1:
                doc.add_page_break()

        total1 = round2(rp1.get('total_in_category', 0), 2)
        total2 = round2(rp2.get('total_in_category', 0), 2)
        total_diff = round2(diff.get('total_in_category', 0), 2)

        total_table = doc.add_table(rows=1, cols=total_cols)
        total_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        c_t0 = total_table.cell(0, 0)
        c_t0.text = _('Total')
        _style_table_cell(c_t0, bold=True, font_size=8)
        c_t1 = total_table.cell(0, 1)
        c_t1.text = str(total1)
        _style_table_cell(c_t1, bold=True, font_size=8)
        c_t2 = total_table.cell(0, 2)
        c_t2.text = str(total2)
        _style_table_cell(c_t2, bold=True, font_size=8)
        c_t3 = total_table.cell(0, 3)
        c_t3.text = str(total_diff)
        _style_table_cell(c_t3, bold=True, font_size=8)


def export(report,
           name1,
           name2,
           reporting_start_datetime_local,
           reporting_end_datetime_local,
           period_type,
           language):
    """
    Export report data to DOCX and return base64 encoded string.
    This function maintains the same interface as the PDF exporter.
    Parameter order: report, name1, name2, reporting_start, reporting_end, period_type, language
    """
    exporter = MeterComparisonDOCXExporter(language)
    return exporter.export(report, name1, name2,
                           reporting_start_datetime_local,
                           reporting_end_datetime_local,
                           period_type,
                           language)
