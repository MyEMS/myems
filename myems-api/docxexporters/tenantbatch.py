"""
Tenant Batch DOCX Exporter

This module provides functionality to export tenant batch data to DOCX format.
It generates a comparison table showing energy consumption, carbon emission and
cost for all tenants of a space within a reporting period.

The exported DOCX file includes:
- Cover page with report metadata
- Batch data table (Name, Space, one column per energy category,
  Carbon Emissions, Costs)

Note: the content mirrors the PDF exporter (pdfexporters/tenantbatch.py)
exactly - the same columns, the same rows and the same 2-decimal precision.
The batch report is a wide matrix and has no chart, matching the PDF output.
"""

import base64
import os
import time
import uuid

from decimal import Decimal
from typing import Optional, Dict, Any
import logging

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from core.utilities import get_translation, round2

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

_font_setup_done = False


def setup_chinese_fonts():
    global _font_setup_done
    if _font_setup_done:
        return True

    font_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             '..', 'pdfexporters', 'fonts', 'NotoSansCJK-Regular.ttc')
    try:
        import matplotlib.font_manager as fm
        fm.fontManager.addfont(font_path)
        prop = fm.FontProperties(fname=font_path)
        font_name = prop.get_name()
        plt.rcParams['font.sans-serif'] = [font_name]
        plt.rcParams['axes.unicode_minus'] = False
        logger.info("Successfully loaded bundled font: %s from %s", font_name, font_path)
        _font_setup_done = True
        return True
    except Exception as e:
        logger.warning("Failed to load bundled font from %s: %s", font_path, e)

    plt.rcParams['font.sans-serif'] = ['DejaVu Sans']
    plt.rcParams['axes.unicode_minus'] = False
    logger.warning("Failed to load bundled NotoSansCJK font, using DejaVu Sans")
    return False


def _convert_decimals(obj):
    if isinstance(obj, Decimal):
        return float(obj)
    elif isinstance(obj, dict):
        return {k: _convert_decimals(v) for k, v in obj.items()}
    elif isinstance(obj, list):
        return [_convert_decimals(item) for item in obj]
    elif isinstance(obj, tuple):
        return tuple(_convert_decimals(item) for item in obj)
    return obj


def _set_min_row_height(cell, height_twips=150, exact=False):
    tr = cell._tc.getparent()
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(height_twips))
    trHeight.set(qn('w:hRule'), 'exact' if exact else 'atLeast')
    existing = trPr.find(qn('w:trHeight'))
    if existing is not None:
        trPr.remove(existing)
    trPr.append(trHeight)


def _reset_cell_paragraph_spacing(cell, font_size=9):
    line_twips = max(int(font_size * 20 * 1.15), 120)
    for p in cell.paragraphs:
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pPr = p._p.get_or_add_pPr()
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:before'), '0')
        spacing.set(qn('w:after'), '0')
        spacing.set(qn('w:line'), str(line_twips))
        spacing.set(qn('w:lineRule'), 'exact')
        existing = pPr.find(qn('w:spacing'))
        if existing is not None:
            pPr.remove(existing)
        pPr.append(spacing)
        ind = pPr.find(qn('w:ind'))
        if ind is not None:
            pPr.remove(ind)


def _set_cell_margins_zero(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side in ['top', 'start', 'bottom', 'end']:
        m = OxmlElement('w:' + side)
        m.set(qn('w:w'), '0')
        m.set(qn('w:type'), 'dxa')
        tcMar.append(m)
    existing = tcPr.find(qn('w:tcMar'))
    if existing is not None:
        tcPr.remove(existing)
    tcPr.append(tcMar)


def _remove_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement('w:' + border_name)
        border.set(qn('w:val'), 'none')
        border.set(qn('w:sz'), '0')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'auto')
        tblBorders.append(border)
    existing = tblPr.find(qn('w:tblBorders'))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(tblBorders)
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)


def _style_table_cell(cell, is_header=False, is_green=False, bold=False, font_size=9):
    if not cell.paragraphs[0].runs:
        cell.paragraphs[0].add_run('')
    cell.paragraphs[0].runs[0].font.bold = bold
    cell.paragraphs[0].runs[0].font.size = Pt(font_size)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _reset_cell_paragraph_spacing(cell, font_size=font_size)
    _set_cell_margins_zero(cell)
    row_h = max(int(font_size * 20 * 1.3), 160)
    _set_min_row_height(cell, height_twips=row_h, exact=False)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement('w:' + border_name)
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '666666')
        tcBorders.append(border)
    tcPr.append(tcBorders)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    if is_header:
        shd.set(qn('w:fill'), 'D9E2F3')
    elif is_green:
        shd.set(qn('w:fill'), '90EE90')
    else:
        return
    tcPr.append(shd)


class TenantBatchDOCXExporter:
    """
    Export tenant batch data to DOCX format.
    Generates a wide matrix table (rows = tenants, columns = energy categories)
    matching the PDF exporter layout.
    """

    def __init__(self, language: str = 'zh_CN'):
        font_setup_success = setup_chinese_fonts()
        if not font_setup_success:
            logger.warning("Chinese font setup failed, some text may not display correctly")

        self.language = language
        self.trans = get_translation(language)
        self._ = self.trans.gettext

        self.dpi = 120

    def export(self,
               report: Dict[str, Any],
               space_name: str,
               reporting_start_datetime_local: str,
               reporting_end_datetime_local: str,
               language: str) -> Optional[str]:
        if report is None:
            return None
        start_time = time.time()
        logger.info("Starting DOCX generation for tenant batch")

        docx_filename = self.generate_docx(
            report,
            space_name,
            reporting_start_datetime_local,
            reporting_end_datetime_local,
            language
        )

        result = ''
        if docx_filename and os.path.exists(docx_filename):
            try:
                with open(docx_filename, 'rb') as binary_file:
                    binary_data = binary_file.read()
                result = base64.b64encode(binary_data).decode('utf-8')
            except Exception as e:
                logger.error("Failed to encode DOCX: %s", str(e))
            finally:
                try:
                    os.remove(docx_filename)
                except Exception:
                    pass
        elapsed = time.time() - start_time
        logger.info("DOCX generation completed in %.2fs", elapsed)
        return result

    def generate_docx(self,
                      report: Dict[str, Any],
                      space_name: str,
                      reporting_start_datetime_local: str,
                      reporting_end_datetime_local: str,
                      language: str) -> Optional[str]:
        _ = self._

        filename = str(uuid.uuid4()) + '.docx'

        self.report = _convert_decimals(report)
        self.space_name = space_name
        self.reporting_start = reporting_start_datetime_local
        self.reporting_end = reporting_end_datetime_local

        doc = Document()
        section = doc.sections[0]
        section.orientation = 1
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)

        tenants = self.report.get('tenants', [])
        has_next = isinstance(tenants, list) and len(tenants) > 0

        self._add_cover_page(doc, has_next=has_next)
        self._add_batch_data_table(doc)

        doc.save(filename)
        logger.info("DOCX generated: %s", filename)
        return filename

    def _add_heading_styled(self, doc, text, level=1):
        heading = doc.add_heading(level=level)
        run = heading.add_run(text)
        run.font.bold = True
        run.font.name = 'Arial'
        r = run._element
        r.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
        return heading

    def _add_cover_page(self, doc, has_next=True):
        _ = self._
        img_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                '..', 'excelexporters', 'myems.png')
        if os.path.exists(img_path):
            try:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(img_path, width=Inches(7.0))
            except Exception as e:
                logger.warning("Failed to load logo image: %s", e)

        for _unused in range(2):
            doc.add_paragraph('')

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(_('Tenant Data') + ' - ' + _('Batch Analysis'))
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.name = 'Arial'
        r = run._element
        r.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

        for _unused in range(2):
            doc.add_paragraph('')

        info_data = [
            [_('Space') + ':', self.space_name],
            [_('Reporting Start Datetime') + ':', self.reporting_start],
            [_('Reporting End Datetime') + ':', self.reporting_end],
        ]

        info_table = doc.add_table(rows=len(info_data), cols=2)
        info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _remove_table_borders(info_table)

        half_w = Inches(3.2)
        for i, (label, value) in enumerate(info_data):
            c_label = info_table.cell(i, 0)
            c_label.width = half_w
            c_label.text = label
            c_label.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_lbl = c_label.paragraphs[0].runs[0]
            r_lbl.bold = True
            r_lbl.font.size = Pt(13)
            r_lbl.font.name = 'Arial'
            r_lbl._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

            c_value = info_table.cell(i, 1)
            c_value.width = half_w
            c_value.text = str(value) if value is not None else ''
            c_value.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_val = c_value.paragraphs[0].runs[0]
            r_val.font.size = Pt(13)
            r_val.font.name = 'Arial'
            r_val._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

        if has_next:
            doc.add_page_break()

    def _add_batch_data_table(self, doc):
        """Add the main batch data table showing all tenants with their
        consumption per energy category, carbon emissions and costs.
        Mirrors the PDF exporter columns:
        [Name, Space, one column per energy category, Carbon Emissions, Costs].
        Paginates if there are too many tenant rows.
        """
        _ = self._
        tenants = self.report.get('tenants', [])

        if not tenants:
            return

        energy_categories = self.report.get('energycategories', [])
        n_categories = len(energy_categories)

        col_headers = [_('Name'), _('Space')]
        for category in energy_categories:
            col_headers.append(str(category.get('name', '')) +
                               ' (' + str(category.get('unit_of_measure', '')) + ')')
        col_headers.append(_('Carbon Emissions') + ' (KGCO2E)')
        col_headers.append(_('Costs') + ' (CNY)')

        total_cols = len(col_headers)
        num_tenants = len(tenants)
        rows_per_page = 20

        font_size = max(6, min(9, 72 // total_cols * 2))

        num_pages = (num_tenants + rows_per_page - 1) // rows_per_page

        for page_idx in range(num_pages):
            start_idx = page_idx * rows_per_page
            end_idx = min(start_idx + rows_per_page, num_tenants)
            page_tenants = tenants[start_idx:end_idx]

            if page_idx == 0:
                self._add_heading_styled(
                    doc, self.space_name + ' - ' + _('Batch Analysis'), level=1
                )
            else:
                self._add_heading_styled(
                    doc, self.space_name + ' - ' + _('Batch Analysis') +
                         ' (' + str(page_idx + 1) + '/' + str(num_pages) + ')',
                    level=1
                )

            num_rows = len(page_tenants) + 1
            table = doc.add_table(rows=num_rows, cols=total_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for j, h in enumerate(col_headers):
                cell = table.cell(0, j)
                cell.text = h
                _style_table_cell(cell, is_header=True, bold=True, font_size=font_size)

            for local_i, tenant in enumerate(page_tenants):
                r_idx = local_i + 1
                values = tenant.get('values', []) or []

                c_name = table.cell(r_idx, 0)
                c_name.text = str(tenant.get('tenant_name', ''))
                _style_table_cell(c_name, bold=True, font_size=font_size)

                c_space = table.cell(r_idx, 1)
                c_space.text = str(tenant.get('space_name', ''))
                _style_table_cell(c_space, bold=True, font_size=font_size)

                for j in range(n_categories):
                    col = 2 + j
                    c_val = table.cell(r_idx, col)
                    value = values[j] if j < len(values) else None
                    c_val.text = str(round2(value, 2)) if value is not None else ''
                    _style_table_cell(c_val, font_size=font_size)

                carbon_col = 2 + n_categories
                c_carbon = table.cell(r_idx, carbon_col)
                carbon = tenant.get('carbon_emissions', 0.0)
                c_carbon.text = str(round2(carbon, 2)) if carbon is not None else ''
                _style_table_cell(c_carbon, font_size=font_size)

                cost_col = 2 + n_categories + 1
                c_cost = table.cell(r_idx, cost_col)
                cost = tenant.get('cost', 0.0)
                c_cost.text = str(round2(cost, 2)) if cost is not None else ''
                _style_table_cell(c_cost, font_size=font_size)

            if page_idx < num_pages - 1:
                doc.add_page_break()


def export(report,
           space_name,
           reporting_start_datetime_local,
           reporting_end_datetime_local,
           language):
    """
    Export report data to DOCX and return base64 encoded string.
    This function maintains the same interface as the PDF exporter.
    """
    exporter = TenantBatchDOCXExporter(language)
    return exporter.export(report,
                         space_name,
                         reporting_start_datetime_local,
                         reporting_end_datetime_local,
                         language)
