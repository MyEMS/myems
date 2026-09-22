"""
Power Quality DOCX Exporter

This module provides functionality to export power quality data to DOCX format.
It generates comprehensive reports showing multi-series point trends for meters
plus the power quality analysis table (voltage/current/frequency compliance).

Key Features:
- Power quality trend table (Datetime + one column per point) with line chart
- Power quality analysis table (Point/Category/Type/Unit/Limit/Normal Limit/
  Severe Limit/Compliance/Worst Deviation/Worst Time + dynamic metrics)
- Multi-language support
- Base64 encoding for file transmission

The exported DOCX file includes:
- Cover page with meter and reporting period info
- Trend pages (paginated table + combined line chart)
- Analysis pages (paginated table, only when report['analysis'] is non-empty)

Note: the content mirrors the PDF exporter (pdfexporters/powerquality.py)
exactly - the same columns, the same rows, the same 3-decimal precision for
trend values and raw values for the analysis table.
"""

import base64
import os
import time
import uuid
import io

from decimal import Decimal
from typing import Optional, Dict, Any, BinaryIO
import logging

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from core.utilities import get_translation, round2

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

_font_setup_done = False


def setup_chinese_fonts():
    global _font_setup_done
    if _font_setup_done:
        return True

    font_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             '..', 'pdfexporters', 'fonts', 'NotoSansCJK-Regular.ttc')
    try:
        import matplotlib.font_manager as fm
        fm.fontManager.addfont(font_path)
        prop = fm.FontProperties(fname=font_path)
        font_name = prop.get_name()
        plt.rcParams['font.sans-serif'] = [font_name]
        plt.rcParams['axes.unicode_minus'] = False
        logger.info(f"Successfully loaded bundled font: {font_name} from {font_path}")
        _font_setup_done = True
        return True
    except Exception as e:
        logger.warning(f"Failed to load bundled font from {font_path}: {e}")

    plt.rcParams['font.sans-serif'] = ['DejaVu Sans']
    plt.rcParams['axes.unicode_minus'] = False
    logger.warning("Failed to load bundled NotoSansCJK font, using DejaVu Sans")
    return False


def _convert_decimals(obj):
    if isinstance(obj, Decimal):
        return float(obj)
    elif isinstance(obj, dict):
        return {k: _convert_decimals(v) for k, v in obj.items()}
    elif isinstance(obj, list):
        return [_convert_decimals(item) for item in obj]
    elif isinstance(obj, tuple):
        return tuple(_convert_decimals(item) for item in obj)
    return obj


def _set_min_row_height(cell, height_twips=150, exact=False):
    tr = cell._tc.getparent()
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(height_twips))
    trHeight.set(qn('w:hRule'), 'exact' if exact else 'atLeast')
    existing = trPr.find(qn('w:trHeight'))
    if existing is not None:
        trPr.remove(existing)
    trPr.append(trHeight)


def _reset_cell_paragraph_spacing(cell, font_size=9):
    line_twips = max(int(font_size * 20 * 1.15), 120)
    for p in cell.paragraphs:
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pPr = p._p.get_or_add_pPr()
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:before'), '0')
        spacing.set(qn('w:after'), '0')
        spacing.set(qn('w:line'), str(line_twips))
        spacing.set(qn('w:lineRule'), 'exact')
        existing = pPr.find(qn('w:spacing'))
        if existing is not None:
            pPr.remove(existing)
        pPr.append(spacing)
        ind = pPr.find(qn('w:ind'))
        if ind is not None:
            pPr.remove(ind)


def _set_cell_margins_zero(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side in ['top', 'start', 'bottom', 'end']:
        m = OxmlElement(f'w:{side}')
        m.set(qn('w:w'), '0')
        m.set(qn('w:type'), 'dxa')
        tcMar.append(m)
    existing = tcPr.find(qn('w:tcMar'))
    if existing is not None:
        tcPr.remove(existing)
    tcPr.append(tcMar)


def _remove_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        border.set(qn('w:sz'), '0')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'auto')
        tblBorders.append(border)
    existing = tblPr.find(qn('w:tblBorders'))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(tblBorders)
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)


def _style_table_cell(cell, is_header=False, is_green=False, bold=False, font_size=9):
    cell.paragraphs[0].runs[0].font.bold = bold
    cell.paragraphs[0].runs[0].font.size = Pt(font_size)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _reset_cell_paragraph_spacing(cell, font_size=font_size)
    _set_cell_margins_zero(cell)
    row_h = max(int(font_size * 20 * 1.3), 160)
    _set_min_row_height(cell, height_twips=row_h, exact=False)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '666666')
        tcBorders.append(border)
    tcPr.append(tcBorders)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    if is_header:
        shd.set(qn('w:fill'), 'D9E2F3')
    elif is_green:
        shd.set(qn('w:fill'), '90EE90')
    else:
        return
    tcPr.append(shd)


class PowerQualityDOCXExporter:
    """
    Export power quality data to DOCX format.
    Generates a multi-series trend table with a combined line chart plus the
    power quality analysis table, matching the PDF exporter layout.
    """

    def __init__(self, language: str = 'zh_CN'):
        font_setup_success = setup_chinese_fonts()
        if not font_setup_success:
            logger.warning("Chinese font setup failed, some text may not display correctly")

        self.language = language
        self.trans = get_translation(language)
        self._ = self.trans.gettext

        self.dpi = 120

        self.chart_colors = ['#4472C4', '#ED7D31', '#70AD47', '#FFC000', '#5B9BD5',
                             '#A855F7', '#EC4899', '#14B8A6', '#F97316', '#6366F1']

    def export(self,
               report: Dict[str, Any],
               name: str,
               reporting_start_datetime_local: str,
               reporting_end_datetime_local: str,
               period_type: str,
               language: str) -> Optional[str]:
        if report is None:
            return None
        start_time = time.time()
        logger.info(f"Starting DOCX generation for {name}")

        docx_filename = self.generate_docx(
            report, name,
            reporting_start_datetime_local,
            reporting_end_datetime_local,
            period_type,
            language
        )

        result = ''
        if docx_filename and os.path.exists(docx_filename):
            try:
                with open(docx_filename, 'rb') as binary_file:
                    binary_data = binary_file.read()
                result = base64.b64encode(binary_data).decode('utf-8')
            except Exception as e:
                logger.error(f"Failed to encode DOCX: {str(e)}")
            finally:
                try:
                    os.remove(docx_filename)
                except Exception:
                    pass
        elapsed = time.time() - start_time
        logger.info(f"DOCX generation completed in {elapsed:.2f}s for {name}")
        return result

    @staticmethod
    def _fig_to_bytesio(fig, dpi: int) -> BinaryIO:
        buf = io.BytesIO()
        fig.tight_layout()
        fig.savefig(buf, format='png', dpi=dpi, bbox_inches='tight')
        plt.close(fig)
        buf.seek(0)
        return buf

    def generate_docx(self,
                      report: Dict[str, Any],
                      name: str,
                      reporting_start_datetime_local: str,
                      reporting_end_datetime_local: str,
                      period_type: str,
                      language: str) -> Optional[str]:
        _ = self._

        filename = str(uuid.uuid4()) + '.docx'

        self.report = _convert_decimals(report)
        self.name = name
        self.reporting_start = reporting_start_datetime_local
        self.reporting_end = reporting_end_datetime_local

        doc = Document()
        section = doc.sections[0]
        section.orientation = 1
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)

        reporting_data = self.report.get('reporting_period', {})
        names = reporting_data.get('names', [])
        has_trend = isinstance(names, list) and len(names) > 0
        analysis = self.report.get('analysis')
        has_analysis = isinstance(analysis, list) and len(analysis) > 0
        self._add_cover_page(doc, has_next=(has_trend or has_analysis))
        self._add_trend_pages(doc, has_next=has_analysis)
        self._add_analysis_pages(doc)

        doc.save(filename)
        logger.info(f"DOCX generated: {filename}")
        return filename

    def _add_heading_styled(self, doc, text, level=1):
        heading = doc.add_heading(level=level)
        run = heading.add_run(text)
        run.font.bold = True
        run.font.name = 'Arial'
        r = run._element
        r.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
        return heading

    def _add_cover_page(self, doc, has_next=True):
        _ = self._
        img_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                '..', 'excelexporters', 'myems.png')
        if os.path.exists(img_path):
            try:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(img_path, width=Inches(7.0))
            except Exception as e:
                logger.warning(f"Failed to load logo image: {e}")

        for _unused in range(2):
            doc.add_paragraph('')

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(_('Power Quality Data') + ' - ' + _('Power Quality Analysis'))
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.name = 'Arial'
        r = run._element
        r.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

        for _unused in range(2):
            doc.add_paragraph('')

        info_data = [
            [_('Name') + ':', self.name],
            [_('Reporting Start Datetime') + ':', self.reporting_start],
            [_('Reporting End Datetime') + ':', self.reporting_end],
        ]

        info_table = doc.add_table(rows=len(info_data), cols=2)
        info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _remove_table_borders(info_table)

        half_w = Inches(3.2)
        for i, (label, value) in enumerate(info_data):
            c_label = info_table.cell(i, 0)
            c_label.width = half_w
            c_label.text = label
            c_label.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_lbl = c_label.paragraphs[0].runs[0]
            r_lbl.bold = True
            r_lbl.font.size = Pt(13)
            r_lbl.font.name = 'Arial'
            r_lbl._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

            c_value = info_table.cell(i, 1)
            c_value.width = half_w
            c_value.text = str(value) if value is not None else ''
            c_value.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_val = c_value.paragraphs[0].runs[0]
            r_val.font.size = Pt(13)
            r_val.font.name = 'Arial'
            r_val._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

        if has_next:
            doc.add_page_break()

    def _add_trend_pages(self, doc, has_next=True):
        _ = self._

        reporting_data = self.report.get('reporting_period', {})
        names = reporting_data.get('names', [])
        timestamps = reporting_data.get('timestamps', [])
        values = reporting_data.get('values', [])

        if names is None or len(names) == 0:
            return

        ca_len = len(names)

        time_axis = []
        for ts_list in timestamps:
            if ts_list is not None and len(ts_list) > 0:
                time_axis = ts_list
                break
        if len(time_axis) == 0:
            return

        num_rows = len(time_axis)
        rows_per_page = 25
        num_pages = (num_rows + rows_per_page - 1) // rows_per_page
        marker_step = max(1, rows_per_page // 15)

        col_headers = [_('Datetime')] + [str(names[i]) for i in range(ca_len)]
        num_cols = len(col_headers)

        datetime_w = 0.22
        series_w = (1.0 - datetime_w) / ca_len
        col_widths = [datetime_w] + [series_w] * ca_len

        if num_cols <= 5:
            fontsize = 7
        elif num_cols <= 9:
            fontsize = 6
        else:
            fontsize = 5

        nan = float('nan')

        for page in range(num_pages):
            start_row = page * rows_per_page
            end_row = min(start_row + rows_per_page, num_rows)
            page_len = end_row - start_row

            if page == 0:
                self._add_heading_styled(doc, self.name + ' ' + _('Trend'), level=1)

            container = doc.add_table(rows=2, cols=1)
            container.alignment = WD_TABLE_ALIGNMENT.CENTER
            _remove_table_borders(container)
            table_cell = container.cell(0, 0)
            chart_cell = container.cell(1, 0)

            table_data = [col_headers]
            for j in range(start_row, end_row):
                row = [str(time_axis[j])]
                for i in range(ca_len):
                    cell = ''
                    if i < len(values) and values[i] is not None and \
                            j < len(values[i]) and values[i][j] is not None:
                        cell = str(round2(values[i][j], 3))
                    row.append(cell)
                table_data.append(row)

            data_table = table_cell.add_table(rows=len(table_data), cols=num_cols)
            data_table.alignment = WD_TABLE_ALIGNMENT.CENTER

            tc = table_cell._tc
            p_elements = tc.findall(qn('w:p'))
            for p_elem in p_elements:
                if len(p_elem.findall(qn('w:r'))) == 0:
                    tc.remove(p_elem)

            for j in range(num_cols):
                cell = data_table.cell(0, j)
                cell.text = col_headers[j]
                _style_table_cell(cell, is_green=True, bold=True, font_size=fontsize)

            for i in range(1, len(table_data)):
                for j in range(num_cols):
                    cell = data_table.cell(i, j)
                    cell.text = table_data[i][j]
                    _style_table_cell(cell, font_size=fontsize)

            fig, ax_chart = plt.subplots(figsize=(8.5, 2.8))
            for i in range(ca_len):
                series_vals = []
                for j in range(start_row, end_row):
                    if i < len(values) and values[i] is not None and \
                            j < len(values[i]) and values[i][j] is not None:
                        series_vals.append(float(values[i][j]))
                    else:
                        series_vals.append(nan)
                color = self.chart_colors[i % len(self.chart_colors)]
                ax_chart.plot(range(page_len), series_vals, linewidth=1.2,
                              color=color, marker='o', markersize=2.5,
                              markevery=marker_step, label=str(names[i]))

            step = max(1, page_len // 8)
            ax_chart.set_xticks(range(0, page_len, step))
            ax_chart.set_xticklabels([str(time_axis[start_row + t])
                                      for t in range(0, page_len, step)],
                                     rotation=45, ha='right', fontsize=7)
            ax_chart.set_title(_('Trend'), fontsize=10, weight='bold')
            legend_ncol = 1 if ca_len <= 4 else 2
            ax_chart.legend(fontsize=6, loc='best', ncol=legend_ncol, framealpha=0.6)
            ax_chart.grid(True, alpha=0.3)
            chart_buf = self._fig_to_bytesio(fig, self.dpi)

            p = chart_cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(chart_buf, width=Inches(8.5))

            if page < num_pages - 1:
                doc.add_page_break()

        if has_next:
            doc.add_page_break()

    def _add_analysis_pages(self, doc):
        _ = self._

        analysis = self.report.get('analysis')
        if not isinstance(analysis, list) or len(analysis) == 0:
            return

        metrics_names = []
        for item in analysis:
            if isinstance(item.get('metrics'), list):
                for m in item['metrics']:
                    if m.get('name') not in metrics_names:
                        metrics_names.append(m.get('name'))
        metrics_names = sorted(metrics_names)

        col_headers = [
            _('Point'), _('Category'), _('Type'), _('Unit'),
            _('Limit'), _('Normal Limit'), _('Severe Limit'), _('Compliance'),
            _('Worst Deviation'), _('Worst Time'),
        ] + metrics_names
        num_cols = len(col_headers)

        all_rows = []
        for item in analysis:
            row_values = [
                item.get('point_name', ''),
                item.get('category', ''),
                item.get('type', ''),
                item.get('unit', ''),
                item.get('limit_pct', ''),
                item.get('limit_normal_hz', ''),
                item.get('limit_severe_hz', ''),
                item.get('compliance_pct', ''),
                item.get('worst_abs_deviation_pct',
                         item.get('worst_unbalance_pct', item.get('worst_deviation_hz', ''))),
                item.get('worst_time', ''),
            ]
            metrics_map = {}
            if isinstance(item.get('metrics'), list):
                for m in item['metrics']:
                    metrics_map[m.get('name')] = m.get('value')
            for mn in metrics_names:
                row_values.append(metrics_map.get(mn, ''))
            all_rows.append(['' if v is None else str(v) for v in row_values])

        num_rows = len(all_rows)
        rows_per_page = 30
        num_pages = (num_rows + rows_per_page - 1) // rows_per_page

        font_size = max(4, min(8, 80 // num_cols))

        for page in range(num_pages):
            start_idx = page * rows_per_page
            end_idx = min(start_idx + rows_per_page, num_rows)

            if num_pages == 1:
                self._add_heading_styled(doc, self.name + ' ' + _('Power Quality Analysis'), level=1)
            else:
                self._add_heading_styled(doc, self.name + ' ' + _('Power Quality Analysis') +
                                         ' (' + str(page + 1) + '/' + str(num_pages) + ')', level=1)

            table_data = [col_headers] + all_rows[start_idx:end_idx]
            data_table = doc.add_table(rows=len(table_data), cols=num_cols)
            data_table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for j in range(num_cols):
                cell = data_table.cell(0, j)
                cell.text = col_headers[j]
                _style_table_cell(cell, is_green=True, bold=True, font_size=font_size)

            for i in range(1, len(table_data)):
                for j in range(num_cols):
                    cell = data_table.cell(i, j)
                    cell.text = table_data[i][j]
                    _style_table_cell(cell, font_size=font_size)

            if page < num_pages - 1:
                doc.add_page_break()


def export(report, name, reporting_start_datetime_local, reporting_end_datetime_local,
           period_type, language):
    """
    Export report data to DOCX and return base64 encoded string.
    This function maintains the same interface as the Excel exporter.
    """
    exporter = PowerQualityDOCXExporter(language)
    return exporter.export(report, name, reporting_start_datetime_local,
                           reporting_end_datetime_local, period_type, language)
