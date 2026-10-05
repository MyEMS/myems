"""
MyEMS DOCX Exporter Common Utilities

This module hosts reusable, report-agnostic helpers used by every DOCX exporter
under the ``docxexporters`` package. It is intentionally kept small so it can
grow organically – any logic that is not specific to a single exporter should
eventually live here.

Currently provided:

* :func:`configure_cover_section` – strips a cover-page Section of every
  visible header/footer content and reduces header/footer distance to zero
  so the cover layout is not affected.
* :func:`configure_body_section`   – attaches the standard MyEMS header
  (centered title + light-grey bottom separator) and a centered page-number
  field to a post-cover Section, resetting the page counter to 1 so cover
  pages do not contribute to the body page numbering.

Additional helpers (shared chart builders, styling primitives, cover page
templates, etc.) can be added here in the future.
"""

from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ---------------------------------------------------------------------------
# Section header / footer configuration
# ---------------------------------------------------------------------------
def configure_cover_section(section):
    """Strip a cover-page ``Section`` of any header or footer content.

    Steps performed:
        1. Unlink header/footer from the previous section so we don't inherit
           anything from a template default.
        2. Clear every run and paragraph text in both header and footer, and
           force paragraph spacing to 0 so even empty paragraphs don't add
           vertical spacing.
        3. Set ``header_distance`` and ``footer_distance`` to 0 so the cover
           content area extends as close to the page edges as the normal
           margins allow.
    """
    section.header.is_linked_to_previous = False
    for p in section.header.paragraphs:
        for r in list(p.runs):
            if r._r.getparent() is not None:
                r._r.getparent().remove(r._r)
        p.text = ''
        pPr = p._p.get_or_add_pPr()
        spacing = pPr.find(qn('w:spacing'))
        if spacing is None:
            spacing = OxmlElement('w:spacing')
            pPr.append(spacing)
        spacing.set(qn('w:before'), '0')
        spacing.set(qn('w:after'), '0')

    section.footer.is_linked_to_previous = False
    for p in section.footer.paragraphs:
        for r in list(p.runs):
            if r._r.getparent() is not None:
                r._r.getparent().remove(r._r)
        p.text = ''
        pPr = p._p.get_or_add_pPr()
        spacing = pPr.find(qn('w:spacing'))
        if spacing is None:
            spacing = OxmlElement('w:spacing')
            pPr.append(spacing)
        spacing.set(qn('w:before'), '0')
        spacing.set(qn('w:after'), '0')

    section.header_distance = Inches(0.0)
    section.footer_distance = Inches(0.0)


def configure_body_section(section, header_title: str = ''):
    """Attach a standard MyEMS DOCX header + footer to a post-cover ``Section``.

    Behavior:
        * **Header** – a single centered line showing ``header_title`` (font
          size 9pt, Arial / SimSun, light-grey bottom border). If
          ``header_title`` is empty the header is still created, just without
          text content.
        * **Footer** – a single centered ``PAGE`` field so Word computes the
          page number automatically. The displayed placeholder is ``"1"``
          until the document is opened/printed.
        * **Page counter** – ``<w:pgNumType w:start="1"/>`` is injected into
          the section properties so the page numbering restarts at 1 for this
          section (i.e. cover pages do not contribute to the body counter).
        * **Distance from edges** – ``header_distance`` and ``footer_distance``
          are both set to ``0.2`` inches so the header sits closer to the top
          edge of the paper, and the footer closer to the bottom edge, giving
          the document body more breathing room.

    Parameters
    ----------
    section:
        Any ``docx.section.Section`` (typically ``doc.sections[1]`` when the
        cover page occupies section 0).
    header_title:
        Pre-formatted text to show in the centered header line. Callers are
        expected to perform any i18n lookups and concatenations themselves.
    """
    # --- Geometry ---------------------------------------------------------
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)

    # --- Page number restart: body pages always start at 1 ---------------
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    pgNumType.set(qn('w:start'), '1')

    # --- Header -----------------------------------------------------------
    header = section.header
    header.is_linked_to_previous = False
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if header_title:
        run = header_para.add_run(header_title)
        run.font.size = Pt(9)
        run.font.name = 'Arial'
        r = run._element
        r.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
    pPr = header_para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '999999')
    pBdr.append(bottom)
    existing_bdr = pPr.find(qn('w:pBdr'))
    if existing_bdr is not None:
        pPr.remove(existing_bdr)
    pPr.append(pBdr)

    # --- Footer (centered page number field) -----------------------------
    footer = section.footer
    footer.is_linked_to_previous = False
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run_begin = footer_para.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    run_begin._r.append(fldChar1)

    run_instr = footer_para.add_run()
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE   \\* MERGEFORMAT '
    run_instr._r.append(instrText)

    run_sep = footer_para.add_run()
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    run_sep._r.append(fldChar2)

    run_text = footer_para.add_run('1')
    run_text.font.size = Pt(9)

    run_end = footer_para.add_run()
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run_end._r.append(fldChar3)
